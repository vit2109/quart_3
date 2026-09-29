"""Парсинг текстовых и табличных файлов в фрагменты для векторной БД."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List, Tuple

import polars as pl

from infrastructure.storage.schema_reader import read_dataframe
from core.config import settings

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log"}
DOC_SUFFIXES = {".docx"}
PDF_SUFFIXES = {".pdf"}
TABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xls", ".json", ".parquet"}

SUPPORTED_SUFFIXES = TEXT_SUFFIXES | DOC_SUFFIXES | PDF_SUFFIXES | TABLE_SUFFIXES

CHUNK_SIZE = 900
CHUNK_OVERLAP = 120

HEADING_MD_RE = re.compile(r"^(#{1,6})\s+(.+)$")
PAGE_MARKER_RE = re.compile(r"^\[Страница\s+(\d+)\]", re.IGNORECASE)
ROW_MARKER_RE = re.compile(r"^Строка\s+(\d+):", re.IGNORECASE)
NUMBER_RE = re.compile(r"\b\d+(?:[.,]\d+)?\b")


def detect_source_type(suffix: str) -> str:
    """Классифицировать тип источника: text, pdf, docx, tabular."""
    s = suffix.lower()
    if s in TABLE_SUFFIXES:
        return "tabular"
    if s in PDF_SUFFIXES:
        return "pdf"
    if s in DOC_SUFFIXES:
        return "docx"
    return "text"


def parse_file(path: Path) -> Tuple[str, Dict[str, Any]]:
    """Извлечь плоский текст и метаданные из файла."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    meta: Dict[str, Any] = {
        "filename": path.name,
        "suffix": suffix,
        "source_type": detect_source_type(suffix),
        "size_bytes": path.stat().st_size,
    }

    if suffix in TEXT_SUFFIXES:
        max_bytes = int(settings.KNOWLEDGE_MAX_TEXT_BYTES or 2_000_000)
        size = path.stat().st_size
        meta["size_bytes"] = size
        if size > max_bytes:
            raw = path.read_bytes()[:max_bytes]
            text = raw.decode("utf-8", errors="replace")
            meta["truncated"] = True
            meta["truncated_to_bytes"] = max_bytes
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
        meta["encoding"] = "utf-8"
        meta["line_count"] = text.count("\n") + 1
        return text.strip(), meta

    if suffix in DOC_SUFFIXES:
        return _parse_docx(path, meta)

    if suffix in PDF_SUFFIXES:
        return _parse_pdf(path, meta)

    return _parse_tabular(path, meta)


def _parse_docx(path: Path, meta: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    from docx import Document

    doc = Document(str(path))
    parts: List[str] = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    for section in doc.sections:
        for block in (section.header, section.footer):
            for paragraph in block.paragraphs:
                text = paragraph.text.strip()
                if text:
                    parts.append(text)

    meta["paragraph_count"] = len(parts)
    meta["table_count"] = len(doc.tables)
    return "\n\n".join(parts), meta


def _parse_pdf(path: Path, meta: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: List[str] = []
    for i, page in enumerate(reader.pages):
        content = (page.extract_text() or "").strip()
        if content:
            pages.append(f"[Страница {i + 1}]\n{content}")
    meta["page_count"] = len(reader.pages)
    meta["parsed_pages"] = len(pages)
    return "\n\n".join(pages), meta


def _parse_tabular(path: Path, meta: Dict[str, Any]) -> Tuple[str, Dict[str, Any]]:
    df = read_dataframe(path)
    meta["rows"] = df.height
    meta["columns"] = list(df.columns)
    lines: List[str] = []
    for idx, row in enumerate(df.head(5000).to_dicts()):
        parts = [f"{k}: {_fmt(v)}" for k, v in row.items()]
        lines.append(f"Строка {idx + 1}: " + " | ".join(parts))
    if df.height > 5000:
        lines.append(f"... ещё {df.height - 5000} строк не включено в превью")
    return "\n".join(lines), meta


def parse_dataset_rows(dataset_path: Path, *, max_rows: int = 5000) -> Tuple[str, Dict[str, Any]]:
    """Преобразовать загруженный табличный датасет в текст для индексации."""
    df = read_dataframe(dataset_path)
    meta = {
        "source_type": "dataset",
        "rows": df.height,
        "columns": list(df.columns),
        "filename": dataset_path.name,
    }
    lines: List[str] = []
    for idx, row in enumerate(df.head(max_rows).to_dicts()):
        parts = [f"{k}: {_fmt(v)}" for k, v in row.items()]
        lines.append(f"Строка {idx + 1}: " + " | ".join(parts))
    return "\n".join(lines), meta


def _extract_chunk_context(paragraph: str, state: Dict[str, Any]) -> Dict[str, Any]:
    """Обновить контекст раздела/страницы и вернуть метаданные для чанка."""
    first_line = paragraph.split("\n", 1)[0].strip()

    page_match = PAGE_MARKER_RE.match(first_line)
    if page_match:
        state["page_number"] = int(page_match.group(1))

    row_match = ROW_MARKER_RE.match(first_line)
    if row_match:
        state["row_number"] = int(row_match.group(1))

    heading_match = HEADING_MD_RE.match(first_line)
    if heading_match:
        state["section_title"] = heading_match.group(2).strip()
        state["section_level"] = len(heading_match.group(1))
    elif (
        first_line
        and len(first_line) < 120
        and first_line == first_line.upper()
        and re.search(r"[A-ZА-ЯЁ]", first_line)
    ):
        state["section_title"] = first_line

    meta: Dict[str, Any] = {}
    if state.get("section_title"):
        meta["section_title"] = state["section_title"]
    if state.get("section_level") is not None:
        meta["section_level"] = state["section_level"]
    if state.get("page_number") is not None:
        meta["page_number"] = state["page_number"]
    if state.get("row_number") is not None:
        meta["row_number"] = state["row_number"]

    numbers = NUMBER_RE.findall(paragraph)
    if numbers:
        meta["numbers"] = numbers[:25]

    keywords = _extract_keywords(paragraph)
    if keywords:
        meta["keywords"] = keywords[:20]

    return meta


def _extract_keywords(text: str) -> List[str]:
    """Ключевые слова для keyword-поиска (слова длиной > 3 символов)."""
    tokens = re.findall(r"[\w\d_]+", text, re.UNICODE)
    seen: set[str] = set()
    result: List[str] = []
    for tok in tokens:
        low = tok.lower()
        if len(low) <= 3 or low.isdigit() or low in seen:
            continue
        seen.add(low)
        result.append(low)
    return result


def chunk_text(
    text: str,
    *,
    chunk_size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
    base_metadata: Dict[str, Any] | None = None,
) -> List[Dict[str, Any]]:
    """Разбить текст на перекрывающиеся фрагменты с метаданными."""
    text = re.sub(r"\n{3,}", "\n\n", text.strip())
    if not text:
        return []

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        paragraphs = [text]

    merged: List[str] = []
    buf = ""
    for para in paragraphs:
        if len(para) > chunk_size:
            if buf:
                merged.append(buf.strip())
                buf = ""
            merged.extend(_split_long(para, chunk_size, overlap))
            continue
        candidate = f"{buf}\n\n{para}".strip() if buf else para
        if len(candidate) <= chunk_size:
            buf = candidate
        else:
            if buf:
                merged.append(buf.strip())
            buf = para
    if buf:
        merged.append(buf.strip())

    chunks: List[Dict[str, Any]] = []
    context_state: Dict[str, Any] = {}
    for i, body in enumerate(merged):
        meta = dict(base_metadata or {})
        meta.update(_extract_chunk_context(body, context_state))
        meta.update(
            {
                "chunk_index": i,
                "char_count": len(body),
                "word_count": len(body.split()),
            }
        )
        chunks.append({"text": body, "metadata": meta})
    return chunks


def _split_long(text: str, size: int, overlap: int) -> List[str]:
    parts: List[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        parts.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)
    return [p for p in parts if p]


def _fmt(value: Any) -> str:
    if value is None:
        return "—"
    if isinstance(value, float):
        return f"{value:.4g}"
    return str(value)
